"""栞｜関西の喫茶店記録 — サムネイル A/B/C 案 + テンプレート生成スクリプト

写真・タイトル・サブ・装飾線・オーバーレイをすべて別オブジェクトとして
PPTX(1080×1350px 相当 / 4:5) に配置する。文字は画像に焼き込まない。

使い方:
    python3 build.py <元写真.jpg>
出力: out/ 以下に PPTX と素材 PNG
"""
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).parent
OUT = HERE / "out"
ASSETS = OUT / "assets"

W, H = 1080, 1350
PX = 9525  # 1px = 9525 EMU (96dpi)

# ── カラートークン ─────────────────────────────
IVORY = "F3ECDF"
PAPER = "EFE7D8"
BROWN = "3A2A20"   # ダークブラウン（本文・タイトル on ivory）
SHADOW = "24180F"  # オーバーレイ用の深いブラウン（黒ではなく茶）
AMBER = "C98B3E"   # 琥珀
AMBER_LIGHT = "E3B172"
BORDEAUX = "6E2F35"
WHITE = "FFFBF3"   # 純白ではなく、わずかにアイボリー寄りの白

# ── フォント（Canva 上に同名で存在するもの） ──────
F_BOLD = "Zen Old Mincho Black"     # A案：太い明朝
F_TITLE = "Shippori Mincho"         # B案：本文明朝
F_TITLE_XB = "Shippori Mincho ExtraBold"  # C案
F_LATIN = "Cormorant Garamond"


def px(v):
    return Emu(int(round(v * PX)))


def pt_from_px(v):
    return Pt(v * 0.75)


# ── 素材 ──────────────────────────────────────
def crop_photo(src, box_w, box_h, y_anchor, name):
    """元写真を幅いっぱいで使い、指定の縦位置(0=上,1=下)で box 比率に切り出す。
    色・質感には手を加えない（リサイズとトリミングのみ）。"""
    im = Image.open(src).convert("RGB")
    sw, sh = im.size
    ratio = box_w / box_h
    cw, ch = sw, round(sw / ratio)
    if ch > sh:
        ch, cw = sh, round(sh * ratio)
    x0 = (sw - cw) // 2
    y0 = round((sh - ch) * y_anchor)
    out = im.crop((x0, y0, x0 + cw, y0 + ch))
    path = ASSETS / name
    out.save(path, quality=95, subsampling=0)
    return path


def gradient_overlay(w, h, hexcolor, a_top, a_bottom, name, ease=1.6):
    """上→下へ透明になる半透明オーバーレイ（独立した PNG オブジェクト）。"""
    r, g, b = (int(hexcolor[i:i + 2], 16) for i in (0, 2, 4))
    im = Image.new("RGBA", (w, h))
    px_ = im.load()
    for y in range(h):
        t = (y / (h - 1)) ** ease
        a = round(255 * (a_top + (a_bottom - a_top) * t))
        for x in range(w):
            px_[x, y] = (r, g, b, a)
    path = ASSETS / name
    im.save(path)
    return path


# ── PPTX ヘルパー ─────────────────────────────
def name_shape(shape, name):
    shape.name = name
    return shape


def add_rect(slide, x, y, w, h, color, alpha=1.0, name="Shape"):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(color)
    if alpha < 1.0:
        clr = s.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(alpha * 100000)))
    s.line.fill.background()
    s.shadow.inherit = False
    return name_shape(s, name)


def add_line(slide, x1, y1, x2, y2, color, weight_px, name="Line"):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    ln.line.color.rgb = RGBColor.from_string(color)
    ln.line.width = pt_from_px(weight_px)
    return name_shape(ln, name)


def add_text(slide, x, y, w, h, lines, *, font, size, color, spacing=0,
             line_height=1.0, align="l", rotation=0, name="Text",
             latin_font=None, anchor="t"):
    """lines: str か [(text, {override})] のリスト。1行=1段落。
    spacing: 字間(px) / line_height: 行送り倍率"""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        text, ov = (ln, {}) if isinstance(ln, str) else ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[ov.get("align", align)]
        p.line_spacing = ov.get("line_height", line_height)
        if "indent" in ov:  # 段落ごとの字下げ（行のズレでリズムを作る）
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(ov["indent"] * PX)))
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = pt_from_px(ov.get("size", size))
        f.color.rgb = RGBColor.from_string(ov.get("color", color))
        rPr = r._r.get_or_add_rPr()
        sp = ov.get("spacing", spacing)
        if sp:
            rPr.set("spc", str(int(sp * 0.75 * 100)))  # 1/100 pt
        fam = ov.get("font", font)
        lat = ov.get("latin_font", latin_font or fam)
        for tag, face in (("a:latin", lat), ("a:ea", fam), ("a:cs", fam)):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", face)
    if rotation:
        tb.rotation = rotation
    return name_shape(tb, name)


def add_picture(slide, path, x, y, w, h, name):
    pic = slide.shapes.add_picture(str(path), px(x), px(y), px(w), px(h))
    return name_shape(pic, name)


# ── 3案のレイアウト ───────────────────────────
def design_a(slide, photo, ov, c):
    """A案：参考アカウント寄り。太明朝の大見出しを -4° 傾け、一瞬で読ませる。"""
    add_picture(slide, photo, 0, 0, W, H, "01_背景写真")
    add_picture(slide, ov, 0, 0, W, 660, "02_オーバーレイ（上部の陰影）")
    add_text(slide, 92, 78, 900, 44, c["sub"], font=F_TITLE, size=30, color=WHITE,
             spacing=4, name="03_サブタイトル")
    add_text(slide, 70, 128, 960, 330, [
        (c["title"][0], {"size": 112}),
        (c["title"][1], {"size": 156, "indent": 170}),
    ], font=F_BOLD, size=112, color=WHITE, spacing=-4, line_height=1.04,
        rotation=-4, name="04_メインタイトル")
    add_line(slide, 262, 478, 700, 448, AMBER_LIGHT, 6, "05_装飾線（琥珀のアンダーライン）")


def design_b(slide, photo, ov, c):
    """B案：栞の世界観重視。本の頁のようなアイボリーの余白に、明朝の二行。"""
    add_rect(slide, 0, 0, W, H, PAPER, name="01_紙の背景（アイボリー）")
    add_picture(slide, photo, 56, 56, 968, 876, "02_写真")
    add_rect(slide, 56, 56, 968, 876, SHADOW, alpha=0.06, name="03_オーバーレイ（写真の締め）")
    add_text(slide, 56, 978, 968, 36, c["eyebrow"], font=F_LATIN, size=24, color=AMBER,
             spacing=7, align="c", name="04_アイブロウ（欧文）")
    add_text(slide, 56, 1018, 968, 200, c["title"], font=F_TITLE, size=80, color=BROWN,
             spacing=2, line_height=1.16, align="c", name="05_メインタイトル")
    add_line(slide, 218, 1290, 262, 1290, AMBER, 2, "06_装飾線（左）")
    add_line(slide, 818, 1290, 862, 1290, AMBER, 2, "06_装飾線（右）")
    add_text(slide, 56, 1272, 968, 40, c["sub"], font=F_TITLE, size=24, color=BROWN,
             spacing=5, align="c", name="07_サブタイトル")


def design_c(slide, photo, ov, c):
    """C案：中間案。上品な明朝を大きく、二行を段違いに組んで静かな動きを出す。"""
    add_picture(slide, photo, 0, 0, W, H, "01_背景写真")
    add_picture(slide, ov, 0, 0, W, 660, "02_オーバーレイ（上部の陰影）")
    add_line(slide, 92, 101, 140, 101, AMBER_LIGHT, 3, "03_装飾線")
    add_text(slide, 156, 82, 800, 40, c["eyebrow"], font=F_LATIN, size=26, color=WHITE,
             spacing=6, name="04_アイブロウ（欧文）")
    add_text(slide, 80, 136, 960, 280, [
        (c["title"][0], {}),
        (c["title"][1], {"indent": 128}),
    ], font=F_TITLE_XB, size=102, color=WHITE, spacing=-2, line_height=1.14,
        name="05_メインタイトル")
    add_text(slide, 212, 440, 800, 40, c["sub"], font=F_TITLE, size=28, color=WHITE,
             spacing=5, name="06_サブタイトル")


COPY = {
    "A": dict(eyebrow="関西の純喫茶で", title=["絵を描きたくなる", "喫茶店!"],
              sub="いちじくのフルーツサンドと、スケッチブック"),
    "B": dict(eyebrow="KISSATEN NOTES  —  SHIORI", title=["喫茶店で、", "絵を描く午後。"],
              sub="いちじくのフルーツサンドと珈琲"),
    "C": dict(eyebrow="FIG SANDWICH & A SKETCHBOOK", title=["いちじくと、", "スケッチの午後。"],
              sub="関西の喫茶店記録｜栞"),
}
TEMPLATE_COPY = {
    "A": dict(eyebrow="小見出し（場所・状況）", title=["テーマを一言で", "喫茶店!"],
              sub="サブ：主役のメニューや具体名"),
    "B": dict(eyebrow="KISSATEN NOTES  —  SHIORI", title=["一行目の言葉、", "二行目の余韻。"],
              sub="サブ：メニュー名・店名・街"),
    "C": dict(eyebrow="ENGLISH SUBTITLE HERE", title=["主役の名前と、", "過ごし方の午後。"],
              sub="関西の喫茶店記録｜栞"),
}
LAYOUT = {"A": design_a, "B": design_b, "C": design_c}


def new_prs():
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    return prs


def main(src):
    ASSETS.mkdir(parents=True, exist_ok=True)
    full = crop_photo(src, W, H, 0.82, "photo_4x5.jpg")
    framed = crop_photo(src, 968, 876, 1.0, "photo_frame.jpg")
    ov_a = gradient_overlay(W, 660, SHADOW, 0.64, 0.0, "overlay_top_A.png")
    ov_c = gradient_overlay(W, 660, SHADOW, 0.6, 0.0, "overlay_top_C.png")
    photos = {"A": (full, ov_a), "B": (framed, None), "C": (full, ov_c)}

    combined = new_prs()
    blank = combined.slide_layouts[6]
    for key in "ABC":
        LAYOUT[key](combined.slides.add_slide(blank), *photos[key], COPY[key])
        single = new_prs()
        LAYOUT[key](single.slides.add_slide(single.slide_layouts[6]), *photos[key], COPY[key])
        single.save(OUT / f"shiori_thumb_{key}.pptx")
    combined.save(OUT / "shiori_thumbnails_ABC.pptx")

    tpl = new_prs()
    for key in "ABC":
        LAYOUT[key](tpl.slides.add_slide(tpl.slide_layouts[6]), *photos[key], TEMPLATE_COPY[key])
    tpl.save(OUT / "shiori_template.pptx")


if __name__ == "__main__":
    main(sys.argv[1])
